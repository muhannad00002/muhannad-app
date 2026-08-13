"""Excel and CSV export.

The Excel workbook is professional: a single "Oman Wedding Vendors" sheet with
the required columns, a frozen/filterable header, auto-sized columns, and the
Google Maps column rendered as clickable hyperlinks. CSV mirrors the columns
with a UTF-8 BOM so Arabic names open correctly in Excel.
"""

from __future__ import annotations

import csv
import io
from typing import Iterable, List, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import Vendor

# Required export columns (label + attribute).
EXPORT_COLUMNS: List[tuple[str, str]] = [
    ("Business Name", "business_name"),
    ("Category", "category"),
    ("Phone Number", "phone_number"),
    ("Governorate", "governorate"),
    ("Google Maps Location Link", "google_maps_url"),
]

SHEET_NAME = "Oman Wedding Vendors"


def _value(vendor: Vendor, attr: str) -> str:
    return getattr(vendor, attr, None) or ""


def to_excel(vendors: Sequence[Vendor]) -> bytes:
    """Return an .xlsx workbook (as bytes) for the given vendors."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0D9488")

    # Header row.
    for col_idx, (label, _) in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    # Data rows.
    maps_col = next(i for i, (_, a) in enumerate(EXPORT_COLUMNS, start=1)
                    if a == "google_maps_url")
    for row_idx, vendor in enumerate(vendors, start=2):
        for col_idx, (_, attr) in enumerate(EXPORT_COLUMNS, start=1):
            value = _value(vendor, attr)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx == maps_col and value:
                cell.hyperlink = value
                cell.value = "Open in Maps"
                cell.font = Font(color="0563C1", underline="single")

    # Freeze header + enable autofilter.
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}{max(ws.max_row, 1)}"

    # Auto-size columns.
    for col_idx, (label, attr) in enumerate(EXPORT_COLUMNS, start=1):
        width = len(label)
        for vendor in vendors:
            width = max(width, len(_value(vendor, attr)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 4, 60)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_csv(vendors: Sequence[Vendor]) -> bytes:
    """Return CSV bytes (UTF-8 with BOM) for the given vendors."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([label for label, _ in EXPORT_COLUMNS])
    for vendor in vendors:
        writer.writerow([_value(vendor, attr) for _, attr in EXPORT_COLUMNS])
    return ("﻿" + buffer.getvalue()).encode("utf-8")
