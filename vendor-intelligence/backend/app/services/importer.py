"""CSV / Excel / TXT import of manually collected vendors (esp. Instagram).

Parses a supplied file into :class:`InstagramCandidate` objects that flow
through the same enrichment/verification/dedup pipeline as automated discovery.
Recognised columns (case-insensitive): Business Name, Instagram URL, Category,
Phone, Governorate.
"""

from __future__ import annotations

import csv
import io
from typing import List, Optional

from openpyxl import load_workbook

from app.search.instagram import InstagramCandidate, extract_username, parse_manual_lines

_ALIASES = {
    "business_name": {"business name", "business_name", "name"},
    "instagram_url": {"instagram url", "instagram", "instagram_url", "profile"},
    "category": {"category"},
    "phone": {"phone", "phone number", "phone_number"},
    "governorate": {"governorate", "region"},
}


def _header_map(header: List[str]) -> dict:
    idx = {}
    for i, col in enumerate(header):
        key = (col or "").strip().lower()
        for field, names in _ALIASES.items():
            if key in names:
                idx[field] = i
    return idx


def _row_to_candidate(row: List[str], idx: dict) -> Optional[InstagramCandidate]:
    def get(field: str) -> str:
        i = idx.get(field)
        return (row[i].strip() if i is not None and i < len(row) and row[i] else "")

    ig = get("instagram_url")
    name = get("business_name")
    username = extract_username(ig) if ig else None
    if not username and not name:
        return None
    username = username or (name.lower().replace(" ", "_"))
    return InstagramCandidate(
        username=username,
        url=f"https://instagram.com/{username}" if extract_username(ig) else (ig or ""),
        title=name or username,
        snippet=" ".join(filter(None, [get("category"), get("governorate"), get("phone")])),
        source="manual_import",
        source_url=ig or "",
        extra_text=" ".join(filter(None, [name, get("category"), get("governorate"), get("phone")])),
    )


def parse_csv(data: bytes) -> List[InstagramCandidate]:
    """Parse CSV bytes into candidates."""
    text = data.decode("utf-8-sig", errors="replace")
    reader = list(csv.reader(io.StringIO(text)))
    if not reader:
        return []
    idx = _header_map(reader[0])
    if not idx:  # no recognised header -> treat every cell as a possible IG link
        flat = [cell for row in reader for cell in row]
        return parse_manual_lines(flat)
    out: List[InstagramCandidate] = []
    for row in reader[1:]:
        cand = _row_to_candidate(row, idx)
        if cand:
            out.append(cand)
    return out


def parse_excel(data: bytes) -> List[InstagramCandidate]:
    """Parse .xlsx bytes into candidates."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = [[("" if c is None else str(c)) for c in row] for row in ws.iter_rows(values_only=True)]
    if not rows:
        return []
    idx = _header_map(rows[0])
    out: List[InstagramCandidate] = []
    for row in rows[1:]:
        cand = _row_to_candidate(row, idx)
        if cand:
            out.append(cand)
    return out


def parse_txt(data: bytes) -> List[InstagramCandidate]:
    """Parse a plain list of Instagram URLs/handles (one per line)."""
    text = data.decode("utf-8-sig", errors="replace")
    return parse_manual_lines(text.splitlines())


def parse_upload(filename: str, data: bytes) -> List[InstagramCandidate]:
    """Dispatch to the right parser based on file extension."""
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return parse_excel(data)
    if name.endswith(".txt"):
        return parse_txt(data)
    return parse_csv(data)
