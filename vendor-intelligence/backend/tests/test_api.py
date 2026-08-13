"""API + export/import integration tests using FastAPI TestClient."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db import Base, get_db
from app.models import Vendor, VendorStatus
from app.services.export import to_excel
from main import app


@pytest.fixture()
def client():
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool
    )
    Base.metadata.create_all(engine)
    TestingSession = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    def override_get_db():
        db = TestingSession()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db

    # Seed a few vendors.
    db = TestingSession()
    db.add_all([
        Vendor(business_name="Royal Bridal", category="Bridal Boutiques", governorate="Muscat",
               phone_number="+968 9123 4567", google_place_id="P1", instagram_url="https://instagram.com/rb",
               status=VendorStatus.APPROVED.value, google_maps_url="https://maps.google.com/?cid=1"),
        Vendor(business_name="Salalah Flowers", category="Flowers", governorate="Dhofar",
               status=VendorStatus.VERIFIED.value),
        Vendor(business_name="Unknown Shop", category="Gifts", governorate="Muscat",
               status=VendorStatus.NEEDS_VERIFICATION.value),
        Vendor(business_name="Bad Vendor", category="DJ", status=VendorStatus.REJECTED.value),
    ])
    db.commit()
    db.close()

    yield TestClient(app)
    app.dependency_overrides.clear()


def test_health(client):
    assert client.get("/api/health").json()["status"] == "ok"


def test_list_and_filter(client):
    r = client.get("/api/vendors").json()
    assert r["total"] == 4
    muscat = client.get("/api/vendors", params={"governorate": "Muscat"}).json()
    assert muscat["total"] == 2
    has_ig = client.get("/api/vendors", params={"has_instagram": True}).json()
    assert has_ig["total"] == 1


def test_stats(client):
    s = client.get("/api/stats").json()
    assert s["total_vendors"] == 4
    assert s["approved"] == 1
    assert s["verified"] == 1
    assert s["needs_verification"] == 1
    assert s["instagram_accounts"] == 1
    assert s["google_maps_matches"] == 1


def test_update_and_audit(client):
    r = client.patch("/api/vendors/3", json={"status": "Approved", "phone_number": "96899887766"})
    assert r.status_code == 200
    assert r.json()["status"] == "Approved"
    assert r.json()["phone_number"] == "+968 9988 7766"
    audit = client.get("/api/vendors/3/audit").json()
    assert any(a["field"] == "status" for a in audit)


def test_bulk_action(client):
    r = client.post("/api/vendors/bulk", json={"ids": [2, 3], "status": "Approved"})
    assert r.json()["updated"] == 2


def test_default_export_excludes_rejected(client):
    r = client.get("/api/export.xlsx")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.title == "Oman Wedding Vendors"
    names = [ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)]
    assert "Bad Vendor" not in names            # rejected excluded by default
    assert "Royal Bridal" in names              # approved included
    assert ws.freeze_panes == "A2"              # header frozen


def test_export_scope_all(client):
    r = client.get("/api/export.csv", params={"scope": "all"})
    assert r.status_code == 200
    assert "Bad Vendor" in r.content.decode("utf-8-sig")


def test_import_instagram_txt(client):
    file = io.BytesIO(b"https://instagram.com/new_oman_shop\n@another_shop\n")
    r = client.post("/api/import", files={"file": ("list.txt", file, "text/plain")},
                    data={"category": "Flowers", "governorate": "Muscat"})
    assert r.status_code == 200
    body = r.json()
    assert body["parsed"] == 2
    assert body["new_vendors"] == 2


def test_discovery_preview_and_meta(client):
    meta = client.get("/api/discovery/meta").json()
    assert len(meta["categories"]) == 35
    prev = client.get("/api/discovery/preview", params={"category": "Wedding Dresses", "deep": True}).json()
    assert prev["count"] > 5
